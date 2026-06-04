#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script lấy dữ liệu công tơ từ API
Chạy: python3 fetch_meter_data.py
"""

import requests
import json
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ========== CONFIGURATION ==========
API_BASE_URL = "http://14.225.244.63:8899/api"
USERNAME = "verdant"
PASSWORD = "verdantenergy@2026"

# Danh sách 30 công tơ
METERS = [
    {"stt": 1, "meter_no": "2410770895", "site": "VN1-YISHENG", "site_id": "VD016"},
    {"stt": 2, "meter_no": "2411320841", "site": "VN2-RK RESOURCES", "site_id": "VD019"},
    {"stt": 3, "meter_no": "2411320925", "site": "VN2-RK RESOURCES", "site_id": "VD019"},
    {"stt": 4, "meter_no": "2411320957", "site": "VN2-RK RESOURCES", "site_id": "VD019"},
    {"stt": 5, "meter_no": "2410770947", "site": "VN1-DELSON_S1", "site_id": "VD011"},
    {"stt": 6, "meter_no": "2410770496", "site": "VN1-DELSON_S1", "site_id": "VD011"},
    {"stt": 7, "meter_no": "2410770519", "site": "VN1-WEISHENG", "site_id": "VD015"},
    {"stt": 8, "meter_no": "2410770533", "site": "VN1-KINGDOM", "site_id": "VD014"},
    {"stt": 9, "meter_no": "2410770525", "site": "VN1-KINGDOM", "site_id": "VD014"},
    {"stt": 10, "meter_no": "2410770851", "site": "VN1-JIANG MEN", "site_id": "VD013"},
    {"stt": 11, "meter_no": "2410770894", "site": "VN1-GREAT PROCESS", "site_id": "VD012"},
    {"stt": 12, "meter_no": "2411320920", "site": "VN2-MOCAL CREATIVE", "site_id": "VD018"},
    {"stt": 13, "meter_no": "2411320977", "site": "VN2-MOCAL CREATIVE", "site_id": "VD018"},
    {"stt": 14, "meter_no": "2410717617", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 15, "meter_no": "2410770552", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 16, "meter_no": "2410770558", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 17, "meter_no": "2410770588", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 18, "meter_no": "2410770598", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 19, "meter_no": "2410770614", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 20, "meter_no": "2410770628", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 21, "meter_no": "2410770631", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 22, "meter_no": "2410770637", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 23, "meter_no": "2410770640", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 24, "meter_no": "2410770720", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 25, "meter_no": "2410770737", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 26, "meter_no": "2510014176", "site": "VN2-ChengShin Rubber", "site_id": "VD017"},
    {"stt": 27, "meter_no": "2510696851", "site": "VN3-Alena-1A", "site_id": "VD020"},
    {"stt": 28, "meter_no": "2510696860", "site": "VN3-Alena-1A", "site_id": "VD020"},
    {"stt": 29, "meter_no": "2510223457", "site": "VN3-Amara-1A", "site_id": "VD021"},
    {"stt": 30, "meter_no": "2510697069", "site": "VN3-Amara-1A", "site_id": "VD021"},
]

# Ngày lấy dữ liệu (hôm qua)
TARGET_DATE = datetime.now() - timedelta(days=1)
START_DATE = TARGET_DATE.strftime("%Y%m%d000000")
END_DATE = TARGET_DATE.strftime("%Y%m%d235959")

print("=" * 80)
print("🔋 LẤY DỮ LIỆU CÔNG TƠ ĐIỆN")
print("=" * 80)
print(f"Ngày lấy: {TARGET_DATE.strftime('%Y-%m-%d')}")
print(f"Start: {START_DATE}, End: {END_DATE}")
print()

# ========== STEP 1: LOGIN ==========
print("🔐 BƯỚC 1: LOGIN LẤY TOKEN")
print("-" * 80)

token = None
try:
    login_url = f"{API_BASE_URL}/Login?UserAccount={USERNAME}&Password={PASSWORD}"
    print(f"Gọi: {login_url}")
    response = requests.get(login_url, timeout=10)
    login_data = response.json()
    
    if login_data.get('CODE') == '1' or login_data.get('CODE') == 1:
        token = login_data.get('TOKEN')
        print(f"✅ Login thành công!")
        print(f"Token: {token[:30]}...")
    else:
        print(f"❌ Login thất bại: {login_data.get('MESSAGE')}")
        exit(1)
except Exception as e:
    print(f"❌ Lỗi: {str(e)}")
    exit(1)

# ========== STEP 2: FETCH DATA ==========
print()
print("📊 BƯỚC 2: LẤY DỮ LIỆU TỪ {0} CÔNG TƠ".format(len(METERS)))
print("-" * 80)

results = []

for meter_info in METERS:
    meter_no = meter_info['meter_no']
    print(f"[{meter_info['stt']}/30] {meter_info['site']} ({meter_no})...", end=" ", flush=True)
    
    try:
        data_url = f"{API_BASE_URL}/GetMeterDataByDate?MeterNo={meter_no}&StartDate={START_DATE}&EndDate={END_DATE}&Token={token}"
        response = requests.get(data_url, timeout=10)
        data = response.json()
        
        if isinstance(data, list) and len(data) > 0:
            first_record = data[0]
            last_record = data[-1]
            
            # Tính delta
            result = {
                'stt': meter_info['stt'],
                'meter_no': meter_no,
                'site': meter_info['site'],
                'site_id': meter_info['site_id'],
                'records_count': len(data),
                'first_time': first_record.get('DATE_TIME', ''),
                'last_time': last_record.get('DATE_TIME', ''),
            }
            
            # Extract các field cần thiết
            for key in first_record.keys():
                if key not in ['METER_NO', 'DATE_TIME', 'RESERVE1', 'RESERVE2']:
                    try:
                        first_val = float(first_record.get(key, 0))
                        last_val = float(last_record.get(key, 0))
                        delta = last_val - first_val
                        result[f'{key}_delta'] = delta
                        result[f'{key}_first'] = first_val
                        result[f'{key}_last'] = last_val
                    except:
                        pass
            
            results.append(result)
            print(f"✅ ({len(data)} record)")
        else:
            print(f"⚠️ Không có dữ liệu")
            
    except Exception as e:
        print(f"❌ {str(e)[:30]}")

# ========== STEP 3: EXPORT EXCEL ==========
print()
print("💾 BƯỚC 3: XUẤT EXCEL")
print("-" * 80)

if results:
    wb = Workbook()
    ws = wb.active
    ws.title = "Meter Data"
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Headers
    headers = [
        'STT', 'Công tơ', 'Site', 'Site ID', 'Số record',
        'Giờ đầu', 'Giờ cuối',
    ]
    
    # Find all delta fields
    delta_fields = set()
    for result in results:
        for key in result.keys():
            if key.endswith('_delta'):
                field_name = key.replace('_delta', '')
                delta_fields.add(field_name)
    
    delta_fields = sorted(list(delta_fields))
    for field in delta_fields:
        headers.append(f"{field} (Sản lượng)")
    
    # Write headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    
    # Write data
    for row_idx, result in enumerate(results, 2):
        col = 1
        ws.cell(row=row_idx, column=col).value = result['stt']
        col += 1
        ws.cell(row=row_idx, column=col).value = result['meter_no']
        col += 1
        ws.cell(row=row_idx, column=col).value = result['site']
        col += 1
        ws.cell(row=row_idx, column=col).value = result['site_id']
        col += 1
        ws.cell(row=row_idx, column=col).value = result['records_count']
        col += 1
        ws.cell(row=row_idx, column=col).value = result['first_time']
        col += 1
        ws.cell(row=row_idx, column=col).value = result['last_time']
        col += 1
        
        for field in delta_fields:
            value = result.get(f'{field}_delta', '')
            cell = ws.cell(row=row_idx, column=col)
            cell.value = value
            if value != '':
                cell.number_format = '#,##0.00'
            cell.border = border
            col += 1
        
        for col_cell in range(1, col):
            ws.cell(row=row_idx, column=col_cell).border = border
    
    # Set column widths
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 18
    ws.column_dimensions['G'].width = 18
    
    filename = f"meter_data_{TARGET_DATE.strftime('%Y%m%d')}.xlsx"
    wb.save(filename)
    print(f"✅ Lưu file: {filename}")
    print(f"   Dữ liệu: {len(results)} công tơ")
    print()
    
    # Print summary
    print("📈 TÓNG TẮT DỮ LIỆU:")
    print("-" * 80)
    for result in results[:5]:
        print(f"[{result['stt']:2d}] {result['meter_no']} - {result['site']}")
        for field in delta_fields:
            value = result.get(f'{field}_delta', '')
            if value != '':
                print(f"     {field}: {value:,.2f}")
    
    if len(results) > 5:
        print(f"     ... và {len(results) - 5} công tơ khác")
else:
    print("❌ Không lấy được dữ liệu")

print()
print("=" * 80)
print("✅ HOÀN TẤT!")
print("=" * 80)
